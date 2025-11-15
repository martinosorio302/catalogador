"""
Test suite for GET /export/inventory endpoint.

Tests the Excel export functionality to ensure proper file generation,
formatting, and data integrity.
"""

import importlib
import sys
from io import BytesIO

from fastapi.testclient import TestClient

sys.path.insert(
    0,
    str(importlib.util.find_spec("api").loader.path if importlib.util.find_spec("api") else "."),
)
from api.main import app
from engine import trd

# Try to import openpyxl for validation
try:
    import openpyxl
    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False


def test_export_inventory_returns_success():
    """Test that /export/inventory returns successful response."""
    client = TestClient(app)
    response = client.get("/export/inventory")

    if not HAS_OPENPYXL:
        # Should return 501 if openpyxl not available
        assert response.status_code == 501
        assert "not available" in response.json()["detail"].lower()
    else:
        # Should return 200 if openpyxl is available
        assert response.status_code == 200


def test_export_inventory_content_type():
    """Test that /export/inventory returns correct content type."""
    if not HAS_OPENPYXL:
        return  # Skip if openpyxl not available

    client = TestClient(app)
    response = client.get("/export/inventory")

    assert response.status_code == 200
    # Check for Excel MIME type
    assert "spreadsheetml" in response.headers["content-type"]


def test_export_inventory_filename_header():
    """Test that /export/inventory sets correct filename in headers."""
    if not HAS_OPENPYXL:
        return  # Skip if openpyxl not available

    client = TestClient(app)
    response = client.get("/export/inventory")

    assert response.status_code == 200
    content_disposition = response.headers.get("content-disposition", "")
    assert "attachment" in content_disposition
    assert "inventario_trd.xlsx" in content_disposition


def test_export_inventory_file_is_valid_excel():
    """Test that /export/inventory returns valid Excel file."""
    if not HAS_OPENPYXL:
        return  # Skip if openpyxl not available

    client = TestClient(app)
    response = client.get("/export/inventory")

    assert response.status_code == 200

    # Try to load the response content as Excel
    excel_data = BytesIO(response.content)
    wb = openpyxl.load_workbook(excel_data)

    # Should have at least one worksheet
    assert len(wb.worksheets) > 0


def test_export_inventory_worksheet_name():
    """Test that Excel worksheet has correct name."""
    if not HAS_OPENPYXL:
        return  # Skip if openpyxl not available

    client = TestClient(app)
    response = client.get("/export/inventory")
    excel_data = BytesIO(response.content)
    wb = openpyxl.load_workbook(excel_data)
    ws = wb.active

    assert ws.title == "Inventario TRD"


def test_export_inventory_headers_present():
    """Test that Excel file contains expected headers."""
    if not HAS_OPENPYXL:
        return  # Skip if openpyxl not available

    client = TestClient(app)
    response = client.get("/export/inventory")
    excel_data = BytesIO(response.content)
    wb = openpyxl.load_workbook(excel_data)
    ws = wb.active

    # Expected headers (row 1)
    expected_headers = [
        "Código", "Título Serie", "Asunto", "Plazo (años)",
        "Temporalidad", "Destino", "AG", "AP", "OAA", "Total"
    ]

    for col_num, expected_header in enumerate(expected_headers, 1):
        actual_header = ws.cell(row=1, column=col_num).value
        assert actual_header == expected_header


def test_export_inventory_header_styling():
    """Test that headers are styled correctly (blue background, white text, bold)."""
    if not HAS_OPENPYXL:
        return  # Skip if openpyxl not available

    client = TestClient(app)
    response = client.get("/export/inventory")
    excel_data = BytesIO(response.content)
    wb = openpyxl.load_workbook(excel_data)
    ws = wb.active

    # Check first header cell styling
    header_cell = ws.cell(row=1, column=1)

    # Check fill color (openpyxl prefixes with 00 or FF)
    assert header_cell.fill.start_color.rgb in (
        "FF366092", "00366092", "366092"
    )

    # Check font is bold and white
    assert header_cell.font.bold is True
    assert header_cell.font.color.rgb in (
        "FFFFFFFF", "00FFFFFF", "FFFFFF"
    )


def test_export_inventory_row_count():
    """Test that Excel file contains correct number of data rows."""
    if not HAS_OPENPYXL:
        return  # Skip if openpyxl not available

    client = TestClient(app)
    response = client.get("/export/inventory")
    excel_data = BytesIO(response.content)
    wb = openpyxl.load_workbook(excel_data)
    ws = wb.active

    # Count non-empty rows (header + data)
    non_empty_rows = sum(1 for row in ws.iter_rows() if any(cell.value for cell in row))

    # Should have header + TRD_TABLA entries
    expected_rows = 1 + len(trd.TRD_TABLA)

    assert non_empty_rows == expected_rows


def test_export_inventory_data_integrity():
    """Test that exported data matches TRD_TABLA content."""
    if not HAS_OPENPYXL:
        return  # Skip if openpyxl not available

    client = TestClient(app)
    response = client.get("/export/inventory")
    excel_data = BytesIO(response.content)
    wb = openpyxl.load_workbook(excel_data)
    ws = wb.active

    # Check first data row (row 2)
    if len(trd.TRD_TABLA) > 0:
        first_entry = trd.TRD_TABLA[0]

        # Column 1: Código
        assert ws.cell(row=2, column=1).value == first_entry.get("code", "")

        # Column 2: Título Serie
        assert ws.cell(row=2, column=2).value == first_entry.get("titulo", "")

        # Column 3: Asunto
        assert ws.cell(row=2, column=3).value == first_entry.get("asunto", "")


def test_export_inventory_column_widths():
    """Test that columns have adjusted widths (not default)."""
    if not HAS_OPENPYXL:
        return  # Skip if openpyxl not available

    client = TestClient(app)
    response = client.get("/export/inventory")
    excel_data = BytesIO(response.content)
    wb = openpyxl.load_workbook(excel_data)
    ws = wb.active

    # Check that at least one column has non-default width
    # Default width is typically around 8.43
    has_adjusted_width = False
    for column in ws.column_dimensions.values():
        if column.width and column.width > 10:
            has_adjusted_width = True
            break

    assert has_adjusted_width, "No columns have adjusted widths"


def test_export_inventory_performance():
    """Test that /export/inventory responds within acceptable time (< 2s)."""
    if not HAS_OPENPYXL:
        return  # Skip if openpyxl not available

    import time

    client = TestClient(app)

    start = time.time()
    response = client.get("/export/inventory")
    elapsed = time.time() - start

    assert response.status_code == 200
    # Should respond within 2 seconds even for large TRD tables
    assert elapsed < 2.0, f"Export took {elapsed:.3f}s, expected < 2.0s"


def test_export_inventory_no_auth_required():
    """Test that /export/inventory does not require authentication."""
    if not HAS_OPENPYXL:
        return  # Skip if openpyxl not available

    client = TestClient(app)

    # Call without any headers
    response = client.get("/export/inventory")

    # Should work without auth
    assert response.status_code == 200


def test_export_inventory_idempotency():
    """Test that multiple calls to /export/inventory produce consistent structure."""
    if not HAS_OPENPYXL:
        return  # Skip if openpyxl not available

    client = TestClient(app)

    response1 = client.get("/export/inventory")
    response2 = client.get("/export/inventory")

    assert response1.status_code == 200
    assert response2.status_code == 200

    # Parse both as Excel and compare structure/data (not binary content)
    # Binary content may differ due to timestamps in Excel metadata
    wb1 = openpyxl.load_workbook(BytesIO(response1.content))
    wb2 = openpyxl.load_workbook(BytesIO(response2.content))

    ws1 = wb1.active
    ws2 = wb2.active

    # Compare row counts
    assert ws1.max_row == ws2.max_row
    assert ws1.max_column == ws2.max_column

    # Compare data in first 5 rows (header + sample data)
    for row_num in range(1, min(6, ws1.max_row + 1)):
        for col_num in range(1, ws1.max_column + 1):
            val1 = ws1.cell(row=row_num, column=col_num).value
            val2 = ws2.cell(row=row_num, column=col_num).value
            assert val1 == val2
