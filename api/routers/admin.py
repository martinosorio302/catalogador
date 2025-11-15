import logging
import os
from io import BytesIO

from fastapi import APIRouter, Header, HTTPException, Request
from fastapi.responses import StreamingResponse

from engine import trd

try:
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    HAS_OPENPYXL = True
except ImportError:
    HAS_OPENPYXL = False

router = APIRouter()
logger = logging.getLogger("catalogador.admin")


@router.post("/reload")
def reload_data(request: Request, x_admin_token: str | None = Header(None)):
    """Reload runtime data (TRD JSON) from disk.

    Security:
    - Requires X-Admin-Token header if ADMIN_RELOAD_TOKEN env var is set
    - Logs all reload attempts with client IP
    - Recommended: Set ADMIN_RELOAD_TOKEN in production

    If the environment variable `ADMIN_RELOAD_TOKEN` is set, require the exact
    token to be provided in the `X-Admin-Token` request header. If not set,
    the endpoint will allow unauthenticated reloads (useful for controlled
    internal networks)—but it's recommended to set a token in production.
    """
    client_ip = request.client.host if request.client else "unknown"

    # Optional auth by env var
    token_required = os.environ.get("ADMIN_RELOAD_TOKEN")
    if token_required:
        if not x_admin_token:
            logger.warning("Reload attempt without token from %s", client_ip)
            raise HTTPException(status_code=401, detail="Missing X-Admin-Token header")
        if x_admin_token != token_required:
            logger.warning("Reload attempt with invalid token from %s", client_ip)
            raise HTTPException(status_code=403, detail="Invalid token")

    try:
        logger.info(
            "Reload requested from %s",
            request.client.host if request.client else "unknown",
        )
        result = trd.reload_trd_data()
        if not result.get("loaded"):
            raise RuntimeError(result.get("errors") or "unknown")
        logger.info("Reloaded TRD data: %s", result.get("counts"))
        return {
            "ok": True,
            "counts": result.get("counts", {}),
        }
    except (OSError, RuntimeError, KeyError) as e:
        logger.exception("Failed to reload TRD data")
        raise HTTPException(status_code=500, detail=str(e))


@router.get("/export/inventory")
def export_inventory():
    """Export TRD inventory as Excel file."""
    if not HAS_OPENPYXL:
        raise HTTPException(
            status_code=501, detail="Excel export not available (openpyxl not installed)"
        )

    try:
        # Create workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Inventario TRD"

        # Headers
        headers = [
            "Código",
            "Título Serie",
            "Asunto",
            "Plazo (años)",
            "Temporalidad",
            "Destino",
            "AG",
            "AP",
            "OAA",
            "Total",
        ]

        # Style headers
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")

        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col_num)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font

        # Data rows
        for row_num, entry in enumerate(trd.TRD_TABLA, 2):
            ws.cell(row=row_num, column=1, value=entry.get("code", ""))
            ws.cell(row=row_num, column=2, value=entry.get("titulo", ""))
            ws.cell(row=row_num, column=3, value=entry.get("asunto", ""))

            # Parse plazo from valor
            valor = entry.get("valor", "")
            if "años" in valor.lower():
                try:
                    plazo = int("".join(filter(str.isdigit, valor)))
                    ws.cell(row=row_num, column=4, value=plazo)
                except (ValueError, TypeError):
                    ws.cell(row=row_num, column=4, value=valor)
            else:
                ws.cell(row=row_num, column=4, value=valor)

            # Temporalidad based on valor
            if "temporal" in valor.lower():
                temporalidad = "Temporal"
            else:
                temporalidad = "Permanente"
            ws.cell(row=row_num, column=5, value=temporalidad)

            ws.cell(row=row_num, column=6, value=entry.get("destino", ""))
            ws.cell(row=row_num, column=7, value=entry.get("ag", ""))
            ws.cell(row=row_num, column=8, value=entry.get("ap", ""))
            ws.cell(row=row_num, column=9, value=entry.get("oaa", ""))
            ws.cell(row=row_num, column=10, value=entry.get("total", ""))

        # Adjust column widths
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(cell.value)
                except (TypeError, AttributeError):
                    pass
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column_letter].width = adjusted_width

        # Save to BytesIO
        output = BytesIO()
        wb.save(output)
        output.seek(0)

        logger.info("Inventory exported: %d entries", len(trd.TRD_TABLA))

        mime_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        return StreamingResponse(
            output,
            media_type=mime_type,
            headers={"Content-Disposition": ("attachment; filename=inventario_trd.xlsx")},
        )

    except (KeyError, AttributeError, ValueError, OSError) as e:
        logger.exception("Failed to export inventory")
        raise HTTPException(status_code=500, detail=str(e))
